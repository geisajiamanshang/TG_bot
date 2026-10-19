"""远程员工入职协议书：根据 templates/远程员工入职协议书_模板.docx 为每个新人生成一份
填好的入职协议（python-docx 直接编辑模板里"_______________"占位符所在的 run，不重建
文档结构，这样字体/加粗/表格样式都原样保留）。

填写规则（用户上传的步骤说明 + 模板本身）：
- 员工花名 / 岗位名称 / 私人tg / 入职日期：直接取自花名册对应字段。
- 编制组织：取花名册 M 列（org_unit），例如"技术中心"；同一个值也用在"协议签署主体"。
- 序列固定"派驻"，服务单位固定"恒睿"（和 hr_defaults.NEW_HIRE_STATIC_DEFAULTS 的口径一致）。
- 合同到期日 = 入职日期 + 1 年 - 1 天；试用期到期日 = 入职日期 + 2 个月 - 1 天
  （hr_defaults.contract_end_date / probation_end_date）。
- 试用期薪资 / 转正薪资：取花名册 AK / AL 列。
- 底部 SSC 签字固定"扶摇"，SSC主管签字留空，日期是生成协议当天；员工签字/日期留空，
  由员工自己签。
"""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

import docx

from app import hr_defaults

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "远程员工入职协议书_模板.docx"

PLACEHOLDER = "_______________"


def _set_cell_value(cell, new_value: str) -> None:
    """把单元格里值等于 PLACEHOLDER 的那个 run 的文本替换成 new_value，其余 run（比如
    加粗的标签文字）原样不动，这样字体/字号都不会丢。new_value 是空的话就不动——宁可让
    "_______________" 留在文档里提醒人工补，也不要把格子清空看起来像是故意留白。找不到
    占位符 run 就整体追加到最后一个段落，保底不丢数据。"""
    if not new_value:
        return
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text == PLACEHOLDER:
                run.text = new_value
                return
    if cell.paragraphs and cell.paragraphs[-1].runs:
        cell.paragraphs[-1].runs[-1].text += new_value
    elif cell.paragraphs:
        cell.paragraphs[-1].add_run(new_value)


def _set_line_prefixed(cell, prefix: str, new_value: str) -> None:
    """单元格里有一行形如"SSC签字：___________________"，把那一整行替换成
    "{prefix}{new_value}"。用于底部签字区域那种一格多行的情况。"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text.startswith(prefix):
                run.text = f"{prefix}{new_value}"
                return


def build_onboarding_agreement(candidate: dict[str, str], today: date | None = None) -> bytes:
    """返回填好的入职协议 docx 字节内容。candidate 需要包含 chinese_name / org_unit /
    position_title / private_contact / effective_date（或 hire_date）/ trial_salary /
    confirmed_salary 这些字段，缺的字段就留模板原来的占位符（不猜）。"""
    document = docx.Document(str(TEMPLATE_PATH))
    today = today or date.today()

    name = str(candidate.get("chinese_name") or "").strip()
    org_unit = str(candidate.get("org_unit") or "").strip()
    position_title = str(candidate.get("position_title") or "").strip()
    private_contact = str(candidate.get("private_contact") or "").strip()
    effective_date = str(candidate.get("effective_date") or candidate.get("hire_date") or "").strip()
    trial_salary = str(candidate.get("trial_salary") or "").strip()
    confirmed_salary = str(candidate.get("confirmed_salary") or "").strip()

    contract_end = hr_defaults.contract_end_date(effective_date) if effective_date else ""
    probation_end = hr_defaults.probation_end_date(effective_date) if effective_date else ""

    table0 = document.tables[0]
    _set_cell_value(table0.rows[0].cells[1], name)
    _set_cell_value(table0.rows[0].cells[3], org_unit)
    _set_cell_value(table0.rows[1].cells[1], "派驻")
    _set_cell_value(table0.rows[1].cells[3], "恒睿")
    _set_cell_value(table0.rows[2].cells[1], position_title)
    _set_cell_value(table0.rows[2].cells[3], private_contact)
    _set_cell_value(table0.rows[3].cells[1], effective_date)
    _set_cell_value(table0.rows[4].cells[3], contract_end)
    _set_cell_value(table0.rows[5].cells[1], probation_end)
    _set_cell_value(table0.rows[5].cells[3], trial_salary)
    _set_cell_value(table0.rows[6].cells[1], confirmed_salary)

    signature_cell = document.tables[3].rows[0].cells[1]
    _set_line_prefixed(signature_cell, "协议签署主体：", org_unit or "（*例：运营中心）")
    _set_line_prefixed(signature_cell, "SSC签字：", "扶摇")
    _set_line_prefixed(signature_cell, "日期：", today.isoformat())

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def agreement_filename(candidate: dict[str, str]) -> str:
    """花名_入职协议_日期.docx（日期用入职日期，和模版文件名规则一致）。"""
    name = str(candidate.get("chinese_name") or "").strip() or "未知花名"
    effective_date = str(candidate.get("effective_date") or candidate.get("hire_date") or "").strip()
    safe_name = "".join(ch for ch in name if ch not in '\\/:*?"<>|')
    return f"{safe_name}_入职协议_{effective_date}.docx"


def month_folder_name(effective_date: str, today: date | None = None) -> str:
    """协议上传到的 Drive 月份子文件夹名，格式和现有文件夹一致，例如"2026年-10月"。
    优先用入职日期所在的月份；入职日期解析失败就退回今天所在的月份。"""
    try:
        year, month, _day = (int(part) for part in effective_date.split("-"))
        return f"{year}年-{month}月"
    except (ValueError, AttributeError):
        today = today or date.today()
        return f"{today.year}年-{today.month}月"
